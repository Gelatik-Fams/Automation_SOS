"""Compatibility API and CLI/watchdog orchestration.

Engine functions are re-exported under their original names. Keep orchestration,
summary export, and the watchdog timestamp here so existing tes monkeypatches
and the shared export/event state retain their original lookup behavior.
"""

import pandas as pd
import time
import os
import glob
try:
    import gspread
except ImportError:
    gspread = None
from openpyxl import Workbook, load_workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.chart.data_source import AxDataSource, StrRef
from openpyxl.chart.label import DataLabelList
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.utils.cell import quote_sheetname
from watchdog.observers.polling import PollingObserver
from watchdog.events import FileSystemEventHandler

from sos_engine.common import (
    GREEN_BG,
    RED_BG,
    METRIC_LABELS,
    N_METRICS,
    WEEK_ORDER,
    UNKNOWN_SOURCE_DIVISION,
    METADATA_COLS_FOR_DEDUP,
    TARGETS_FILENAME,
    DEFAULT_TARGET,
    TARGET_DIM_ORDER,
    EXCEL_CATEGORY_CHART_ROW_STEP,
    col_letter,
    get_dashboard_division_options,
    filter_dashboard_by_division,
    sort_key_period,
    get_cluster_name,
    get_summary_output_path,
    get_store_detail_output_path,
    get_source_divisions,
)
from sos_engine.calculations import (
    is_competitor,
    extract_parent_brand,
    calc_sos,
    hitung_sos,
)
from sos_engine.data import (
    extract_source_division_from_filename,
    extract_source_division_from_raw_data,
    discover_report_product_files,
    display_csv_path,
    count_files_by_source_division,
    count_physical_text_lines,
    validasi_data,
    baca_semua_csv,
)
from sos_engine.targets import (
    get_target,
    default_target_rows,
    default_targets_dict,
    buat_targets_excel_default,
    _parse_target_value,
    _target_rows_to_dict,
    _build_division_targets,
    _enrich_targets_with_df_values,
    read_targets_from_workbook,
    load_or_create_summary_targets,
    baca_target_dari_excel,
    load_or_create_local_targets,
)
from sos_engine.tables import (
    buat_division_summary_rows,
    buat_tabel_sos_monthly,
    buat_tabel_channel_account,
    buat_category_divisi_section,
    buat_region_divisi_section,
    buat_account_divisi_section,
)
from sos_engine.dashboard import (
    build_dashboard_payload,
    _prune_dashboard_row_for_excel,
    dashboard_payload_sos_only,
    _shift_section_rows,
    build_division_excel_payload,
)
from sos_engine.excel import (
    _excel_color,
    _write_rows,
    sanitize_excel_sheet_name,
    get_excel_conditional_format_ranges,
    _format_excel_dashboard,
    _add_division_summary_chart,
    _excel_text_axis_ref,
    _add_category_by_divisi_charts,
    _add_dashboard_charts,
    _write_targets_sheet,
    _write_validation_report_sheet,
    export_store_detail_excel,
)
from sos_engine.sheets_formatting import (
    clear_dashboard_content,
    terapkan_dropdown_division,
    hapus_semua_chart,
    hapus_semua_formatting,
    terapkan_filter,
    api_retry,
    hapus_conditional_format,
    terapkan_conditional_format,
    tambahkan_chart_category_divisi,
    tambahkan_chart_division_summary,
    sembunyikan_dashboard_row_ranges,
)
from sos_engine.sheets import (
    baca_dashboard_division_selector,
    baca_target_dari_ws_targets,
    baca_target_dari_dashboard,
    inisialisasi_ws_targets,
    buat_dashboard,
    buat_store_detail,
    buat_validation_report,
)


# Shared by export_summary_excel and CSVHandler; keep this state in this module.
_last_excel_write_time = 0.0


def export_summary_excel(df, targets, output_dir='.', cluster_name=None, target_rows=None, df_removed=None):
    divisions = get_source_divisions(df)
    if not divisions:
        raise ValueError('Tidak ada data untuk dashboard.')

    wb = Workbook()
    default_sheet = wb.active
    wb.remove(default_sheet)

    used_sheet_names = set()
    wrote_dashboard = False
    for division in divisions:
        excel_payload = build_division_excel_payload(df, targets, division, target_rows=target_rows)
        if excel_payload is None:
            continue

        ws_dashboard = wb.create_sheet(sanitize_excel_sheet_name(division, used_sheet_names))
        _write_rows(ws_dashboard, excel_payload['rows'])
        _format_excel_dashboard(ws_dashboard, excel_payload)
        _add_dashboard_charts(ws_dashboard, excel_payload)
        wrote_dashboard = True

    if not wrote_dashboard:
        raise ValueError('Tidak ada data untuk dashboard.')

    # Enrich targets to ensure GRAND TOTAL entries exist for every division × dimension
    target_rows = _enrich_targets_with_df_values(df, target_rows)
    if target_rows is not None:
        targets = _target_rows_to_dict(target_rows)

    ws_targets = wb.create_sheet('TARGETS')
    _write_targets_sheet(ws_targets, targets, target_rows=target_rows, divisions=divisions)

    # VALIDATION REPORT sheet
    ws_validation = wb.create_sheet('VALIDATION REPORT')
    _write_validation_report_sheet(ws_validation, df_removed)

    output_path = get_summary_output_path(output_dir, cluster_name, extension='.xlsx')
    wb.save(output_path)
    global _last_excel_write_time
    _last_excel_write_time = time.time()
    return output_path


# CLI orchestration intentionally resolves its dependencies through tes globals.
def proses_data():
    summary_path = get_summary_output_path('.', extension='.xlsx')
    try:
        targets_local, target_rows = load_or_create_summary_targets(summary_path)
    except ValueError as e:
        print(f'[ERROR] {e}')
        return

    print('\nMemproses data...')

    df_raw = baca_semua_csv()
    if df_raw is None:
        return

    df, _df_removed = validasi_data(df_raw)

    target_rows = _enrich_targets_with_df_values(df, target_rows)
    if target_rows is not None:
        targets_local = _target_rows_to_dict(target_rows)

    try:
        output_path = export_summary_excel(
            df,
            targets_local,
            output_dir='.',
            target_rows=target_rows,
            df_removed=_df_removed,
        )
        print(f'Summary SOS berhasil dibuat: {output_path}')
        
        output_store_detail = export_store_detail_excel(
            df,
            output_dir='.'
        )
        print(f'Store Detail berhasil dibuat: {output_store_detail}')
    except ValueError as e:
        print(f'[ERROR] {e}')
        return


class CSVHandler(FileSystemEventHandler):
    def on_any_event(self, event):
        if event.event_type not in ('modified', 'created'):
            return
        path = event.src_path
        if path.endswith('.csv'):
            print(f'File CSV terdeteksi: {path}')
            time.sleep(2)
            proses_data()
        elif (path.endswith('.xlsx')
              and os.path.basename(path).startswith('Summary SOS_')
              and time.time() - _last_excel_write_time > 5):
            print(f'File TARGETS berubah: {path}')
            time.sleep(2)
            proses_data()


if __name__ == '__main__':
    proses_data()  # proses langsung saat pertama dijalankan

    observer = PollingObserver()
    observer.schedule(CSVHandler(), path='.', recursive=True)
    observer.start()
    print('Watchdog aktif! Menunggu perubahan CSV...')

    try:
        while True:
            time.sleep(1)
    except KeyboardInterrupt:
        observer.stop()
    observer.join()
