import unittest
import os
import tempfile

from openpyxl import Workbook, load_workbook
import pandas as pd

import tes


def dataframe_with_produsen(rows):
    df = pd.DataFrame(rows)
    if "Product Code" in df.columns and "Produsen" not in df.columns:
        is_competitor = df["Product Code"].astype(str).str.contains(
            "COMPETITOR", case=False, na=False
        )
        df["Produsen"] = is_competitor.map({True: "COMPETITOR", False: "INDOFOOD"})
    return df


class FakeSpreadsheet:
    def __init__(self):
        self.batch_updates = []

    def fetch_sheet_metadata(self):
        return {"sheets": [{"properties": {"sheetId": 1}}]}

    def batch_update(self, *args, **kwargs):
        self.batch_updates.append((args, kwargs))
        return None


class FakeWorksheet:
    id = 1

    def __init__(self, b1_value=""):
        self.spreadsheet = FakeSpreadsheet()
        self.updated_values = None
        self.frozen = None
        self.b1_value = b1_value
        self.cleared = False
        self.batch_clears = []

    def clear(self):
        self.cleared = True
        return None

    def batch_clear(self, ranges):
        self.batch_clears.append(ranges)

    def update(self, range_name=None, values=None):
        self.updated_values = values

    def acell(self, label):
        class Cell:
            def __init__(self, value):
                self.value = value
        return Cell(self.b1_value if label == "B1" else "")

    def get_all_values(self):
        return []

    def freeze(self, rows=None, cols=None):
        self.frozen = (rows, cols)


class SourceDivisionTests(unittest.TestCase):
    def test_extracts_source_division_from_last_filename_segment(self):
        self.assertEqual(
            tes.extract_source_division_from_filename(
                "Report Product - Mei 25 - Oil & Fat.csv"
            ),
            "Oil & Fat",
        )

    def test_old_filename_without_division_returns_unknown(self):
        self.assertEqual(
            tes.extract_source_division_from_filename(
                "Report Product - Januari 2024.csv"
            ),
            tes.UNKNOWN_SOURCE_DIVISION,
        )

    def test_extracts_source_division_from_raw_divisi_column(self):
        df = pd.DataFrame({"Divisi": ["Pasta", "Pasta", " Pasta "]})

        self.assertEqual(
            tes.extract_source_division_from_raw_data(df, "Report Product - Jan 25.csv"),
            "Pasta",
        )

    def test_extract_source_division_requires_divisi_column(self):
        df = pd.DataFrame({"Visit Date": ["2025-01-02"]})

        with self.assertRaisesRegex(ValueError, "kolom Divisi tidak ditemukan"):
            tes.extract_source_division_from_raw_data(df, "Report Product - Jan 25.csv")

    def test_extract_source_division_rejects_mixed_division_file(self):
        df = pd.DataFrame({"Divisi": ["Pasta", "Noodle"]})

        with self.assertRaisesRegex(ValueError, "lebih dari satu Divisi"):
            tes.extract_source_division_from_raw_data(df, "Report Product - Jan 25.csv")


class FileDiscoveryTests(unittest.TestCase):
    def test_discovers_report_product_csv_in_root_and_division_folders(self):
        with tempfile.TemporaryDirectory() as tmp:
            paths = [
                os.path.join(tmp, "Report Product - Jan 25.csv"),
                os.path.join(tmp, "SIMP", "Report Product - Jan 25 - SIMP.csv"),
                os.path.join(tmp, "NICI", "Report Product - Jan 25 - NICI.csv"),
                os.path.join(tmp, "IFM", "Report Product - Jan 25 - IFM.csv"),
                os.path.join(tmp, "Noodle", "Report Product - Jan 25 - Noodle.csv"),
            ]
            for path in paths:
                os.makedirs(os.path.dirname(path), exist_ok=True)
                with open(path, "w", encoding="utf-8") as f:
                    f.write("Visit Date,Product Code,Facing\n")
            with open(os.path.join(tmp, "Other.csv"), "w", encoding="utf-8") as f:
                f.write("Visit Date,Product Code,Facing\n")

            found = [
                os.path.relpath(p, tmp).replace(os.sep, "/")
                for p in tes.discover_report_product_files(tmp)
            ]

        self.assertEqual(found, [
            "IFM/Report Product - Jan 25 - IFM.csv",
            "NICI/Report Product - Jan 25 - NICI.csv",
            "Noodle/Report Product - Jan 25 - Noodle.csv",
            "Report Product - Jan 25.csv",
            "SIMP/Report Product - Jan 25 - SIMP.csv",
        ])

    def test_counts_files_per_source_division_from_filenames(self):
        files = [
            "SIMP/Report Product - Jan 25 - SIMP.csv",
            "SIMP/Report Product - Feb 25 - SIMP.csv",
            "NICI/Report Product - Jan 25 - NICI.csv",
            "Report Product - Jan 25.csv",
        ]

        self.assertEqual(tes.count_files_by_source_division(files), {
            "NICI": 1,
            "SIMP": 2,
            tes.UNKNOWN_SOURCE_DIVISION: 1,
        })

    def test_baca_semua_csv_keeps_physical_line_count_for_user_log(self):
        original_cwd = os.getcwd()
        with tempfile.TemporaryDirectory() as tmp:
            try:
                os.chdir(tmp)
                with open("Report Product - Jan 25 - Pasta.csv", "w", encoding="utf-8", newline="") as f:
                    f.write("Visit Date,Product Code,Produsen,Divisi,Facing\n")
                    f.write('"2025-01-02","SKU\n001","INDOFOOD","Pasta",10\n')
                    f.write('"2025-01-02","COMPETITOR SKU","COMPETITOR","Pasta",5\n')

                df = tes.baca_semua_csv()

                self.assertEqual(df.attrs["physical_csv_lines"], 4)
                self.assertEqual(len(df), 2)
                self.assertEqual(set(df["Source Division"]), {"Pasta"})
            finally:
                os.chdir(original_cwd)

    def test_baca_semua_csv_uses_divisi_column_even_when_filename_has_no_division(self):
        original_cwd = os.getcwd()
        with tempfile.TemporaryDirectory() as tmp:
            try:
                os.chdir(tmp)
                with open("Report Product - Jan 25.csv", "w", encoding="utf-8", newline="") as f:
                    f.write("Visit Date,Product Code,Produsen,Divisi,Facing\n")
                    f.write('"2025-01-02","SKU001","INDOFOOD","Pasta",10\n')

                df = tes.baca_semua_csv()

                self.assertEqual(df["Source Division"].tolist(), ["Pasta"])
            finally:
                os.chdir(original_cwd)

    def test_baca_semua_csv_reads_each_month_file_and_detects_same_division(self):
        original_cwd = os.getcwd()
        with tempfile.TemporaryDirectory() as tmp:
            try:
                os.chdir(tmp)
                for month in ["Jan", "Feb", "Mar"]:
                    with open(f"Report Product - {month} 25.csv", "w", encoding="utf-8", newline="") as f:
                        f.write("Visit Date,Product Code,Produsen,Divisi,Facing\n")
                        f.write(f'"2025-01-02","SKU-{month}","INDOFOOD","Pasta",10\n')

                df = tes.baca_semua_csv()

                self.assertEqual(len(df), 3)
                self.assertEqual(set(df["Source Division"]), {"Pasta"})
            finally:
                os.chdir(original_cwd)

    def test_baca_semua_csv_fails_when_divisi_column_missing(self):
        original_cwd = os.getcwd()
        with tempfile.TemporaryDirectory() as tmp:
            try:
                os.chdir(tmp)
                with open("Report Product - Jan 25.csv", "w", encoding="utf-8", newline="") as f:
                    f.write("Visit Date,Product Code,Produsen,Facing\n")
                    f.write('"2025-01-02","SKU001","INDOFOOD",10\n')

                with self.assertRaisesRegex(ValueError, "kolom Divisi tidak ditemukan"):
                    tes.baca_semua_csv()
            finally:
                os.chdir(original_cwd)

    def test_baca_semua_csv_fails_when_one_file_has_multiple_divisions(self):
        original_cwd = os.getcwd()
        with tempfile.TemporaryDirectory() as tmp:
            try:
                os.chdir(tmp)
                with open("Report Product - Jan 25.csv", "w", encoding="utf-8", newline="") as f:
                    f.write("Visit Date,Product Code,Produsen,Divisi,Facing\n")
                    f.write('"2025-01-02","SKU001","INDOFOOD","Pasta",10\n')
                    f.write('"2025-01-02","SKU002","INDOFOOD","Noodle",10\n')

                with self.assertRaisesRegex(ValueError, "lebih dari satu Divisi"):
                    tes.baca_semua_csv()
            finally:
                os.chdir(original_cwd)


class AppsScriptRemovalTests(unittest.TestCase):
    def test_apps_script_file_is_not_required_for_excel_output(self):
        self.assertFalse(os.path.exists("dashboard_division_filter.gs"))


class ValidationFreezeTests(unittest.TestCase):
    def test_sos_classifies_competitor_from_produsen_not_product_code(self):
        df = dataframe_with_produsen([
            {
                "Region": "BANDUNG",
                "Period": "Jan 25",
                "Produsen": "INDOFOOD",
                "Product Code": "SKU001",
                "Facing": 30,
            },
            {
                "Region": "BANDUNG",
                "Period": "Jan 25",
                "Produsen": "COMPETITOR",
                "Product Code": "'I401",
                "Facing": 20,
            },
        ])

        result = tes.calc_sos(df, ["Region"])

        self.assertEqual(result.loc[0, "fi"], 30)
        self.assertEqual(result.loc[0, "fk"], 20)
        self.assertEqual(result.loc[0, "SOS%"], 60.0)

    def test_dashboard_does_not_keep_unused_compliance_dead_code(self):
        with open("tes.py", encoding="utf-8") as f:
            source = f.read()

        for token in ("calc_" + "compliance", "compliance_" + "map"):
            self.assertEqual(source.count(token), 0, token)

    def test_source_division_metadata_does_not_change_duplicate_detection(self):
        df = dataframe_with_produsen([
            {
                "_source_file": "Report Product - Jan 25 - Beverage",
                "Source Division": "Beverage",
                "Visit Date": "2025-01-02",
                "Store Code": "S001",
                "Product Code": "SKU001",
                "Facing": 10,
            },
            {
                "_source_file": "Report Product - Jan 25 - Pasta",
                "Source Division": "Pasta",
                "Visit Date": "2025-01-02",
                "Store Code": "S001",
                "Product Code": "SKU001",
                "Facing": 10,
            },
        ])

        cleaned, removed = tes.validasi_data(df)

        self.assertEqual(len(cleaned), 1)
        self.assertEqual(len(removed), 1)


class LocalTargetsFileTests(unittest.TestCase):
    def test_default_target_rows_use_embedded_four_column_structure(self):
        rows = tes.default_target_rows(["Noodle", "Snack"])

        self.assertEqual(rows[0], ["Division", "Dimension", "Name", "Target"])
        self.assertIn(["Noodle", "REGION", "DEFAULT", 65.0], rows)
        self.assertIn(["Snack", "REGION", "DEFAULT", 65.0], rows)

    def test_missing_summary_uses_default_targets_without_creating_targets_file(self):
        with tempfile.TemporaryDirectory() as tmp:
            path = os.path.join(tmp, "Summary SOS_Test.xlsx")

            targets, target_rows = tes.load_or_create_summary_targets(path)

            self.assertFalse(os.path.exists(path))
            self.assertFalse(os.path.exists(os.path.join(tmp, "TARGETS.xlsx")))
            self.assertEqual(targets[("REGION", "DEFAULT")], 65.0)
            self.assertIsNone(target_rows)

    def test_existing_summary_targets_sheet_is_read_without_overwrite(self):
        with tempfile.TemporaryDirectory() as tmp:
            path = os.path.join(tmp, "Summary SOS_Test.xlsx")
            wb = Workbook()
            ws = wb.active
            ws.title = "DASHBOARD"
            targets_ws = wb.create_sheet("TARGETS")
            targets_ws.append(["Division", "Dimension", "Name", "Target"])
            targets_ws.append(["Noodle", "REGION", "WEST", 80])
            targets_ws.append(["Snack", "CHANNEL", "MT", "70,5"])
            wb.save(path)

            targets, target_rows = tes.load_or_create_summary_targets(path)

            self.assertEqual(targets[("REGION", "WEST")], 80.0)
            self.assertEqual(targets[("CHANNEL", "MT")], 70.5)
            self.assertEqual(target_rows[1], ["Noodle", "REGION", "WEST", 80])

            wb_after = load_workbook(path)
            self.assertEqual(wb_after.sheetnames, ["DASHBOARD", "TARGETS"])

    def test_invalid_summary_targets_sheet_raises_clear_error_without_overwrite(self):
        with tempfile.TemporaryDirectory() as tmp:
            path = os.path.join(tmp, "Summary SOS_Test.xlsx")
            wb = Workbook()
            ws = wb.active
            ws.title = "DASHBOARD"
            targets_ws = wb.create_sheet("TARGETS")
            targets_ws.append(["Bad", "Header"])
            wb.save(path)

            with self.assertRaises(ValueError) as ctx:
                tes.load_or_create_summary_targets(path)

            self.assertIn("TARGETS", str(ctx.exception))
            wb_after = load_workbook(path)
            self.assertEqual(wb_after.sheetnames, ["DASHBOARD", "TARGETS"])

    def test_dashboard_target_reader_accepts_local_targets_dict(self):
        ws = FakeWorksheet()

        targets = tes.baca_target_dari_dashboard(ws, {
            ("REGION", "WEST"): 82.0,
        })

        self.assertEqual(targets, {("REGION", "WEST"): 82.0})

    def test_first_run_exports_summary_xlsx_with_embedded_targets_and_no_targets_xlsx(self):
        original_cwd = os.getcwd()
        original_baca_semua_csv = tes.baca_semua_csv
        with tempfile.TemporaryDirectory() as tmp:
            try:
                os.chdir(tmp)

                tes.baca_semua_csv = ExcelExportTests().make_df
                tes.proses_data()

                expected_name = f"Summary SOS_{os.path.basename(tmp)}.xlsx"
                self.assertTrue(os.path.exists(expected_name))
                self.assertFalse(os.path.exists("TARGETS.xlsx"))
                wb = load_workbook(expected_name)
                self.assertEqual(
                    [wb["TARGETS"]["A1"].value, wb["TARGETS"]["B1"].value, wb["TARGETS"]["C1"].value, wb["TARGETS"]["D1"].value],
                    ["Division", "Dimension", "Name", "Target"],
                )
            finally:
                tes.baca_semua_csv = original_baca_semua_csv
                os.chdir(original_cwd)


class ExcelExportTests(unittest.TestCase):
    def make_df(self):
        return DashboardDivisionAwareTests().make_dashboard_df()

    def test_export_summary_excel_creates_only_dashboard_and_targets(self):
        with tempfile.TemporaryDirectory() as tmp:
            output_path = tes.export_summary_excel(
                self.make_df(),
                {("REGION", "WEST"): 70},
                output_dir=tmp,
                cluster_name="Indulgence",
            )

            self.assertEqual(
                os.path.basename(output_path),
                "Summary SOS_Indulgence.xlsx",
            )
            wb = load_workbook(output_path)
            self.assertEqual(wb.sheetnames, ["Noodle", "Snack", "TARGETS"])
            targets_ws = wb["TARGETS"]
            target_divisions = {
                targets_ws.cell(row=row_number, column=1).value
                for row_number in range(2, targets_ws.max_row + 1)
            }
            self.assertEqual(target_divisions, {"Noodle", "Snack"})
            for sheet_name in ["Noodle", "Snack"]:
                ws = wb[sheet_name]
                self.assertEqual(ws["A1"].value, "DIVISION SUMMARY")
                self.assertNotEqual(ws["B1"].value, "ALL")
                self.assertFalse(ws.data_validations.dataValidation)
                self.assertFalse(any(ws.row_dimensions[row].hidden for row in range(1, ws.max_row + 1)))
                values = [row[0] for row in ws.iter_rows(values_only=True) if row and row[0]]
                self.assertFalse(any(str(value).startswith("CHART SOURCE ") for value in values))

    def test_export_summary_excel_creates_one_sheet_per_source_division(self):
        df = self.make_df()
        extra = []
        for division in ["Nici", "Oil & Fat"]:
            row = df.iloc[0].copy()
            row["Source Division"] = division
            row["Region"] = division
            extra.append(row)
            comp = df.iloc[1].copy()
            comp["Source Division"] = division
            comp["Region"] = division
            extra.append(comp)
        df = pd.concat([df, dataframe_with_produsen(extra)], ignore_index=True)

        with tempfile.TemporaryDirectory() as tmp:
            output_path = tes.export_summary_excel(
                df,
                {("REGION", "WEST"): 70},
                output_dir=tmp,
                cluster_name="Indulgence",
            )
            wb = load_workbook(output_path)

        self.assertEqual(wb.sheetnames, ["Nici", "Noodle", "Oil & Fat", "Snack", "TARGETS"])

    def test_excel_dashboard_removes_facing_columns_and_keeps_sos_values(self):
        with tempfile.TemporaryDirectory() as tmp:
            output_path = tes.export_summary_excel(
                self.make_df(),
                {("REGION", "WEST"): 70},
                output_dir=tmp,
                cluster_name="Indulgence",
            )
            ws = load_workbook(output_path)["Noodle"]
            rows = list(ws.iter_rows(values_only=True))

        title_row = next(i for i, row in enumerate(rows) if row[0] == "SOS% BY REGION")
        header1 = list(rows[title_row + 1])
        header2 = list(rows[title_row + 2])
        noodle_row = next(row for row in rows if row[:2] == ("Noodle", "WEST"))

        self.assertIn("TARGET", header1)
        self.assertIn("Jan 25", header1)
        self.assertIn("SOS%", header2)
        self.assertNotIn("Indofood", header2)
        self.assertNotIn("Kompetitor", header2)
        self.assertNotIn("Total", header2)
        self.assertEqual(list(noodle_row[:4]), ["Noodle", "WEST", 70, 75.0])

    def test_excel_conditional_format_ranges_only_target_detail_sos_cells(self):
        payload = tes.dashboard_payload_sos_only(
            tes.build_dashboard_payload(
                self.make_df(),
                {("REGION", "WEST"): 70},
            )
        )

        ranges = tes.get_excel_conditional_format_ranges(payload)

        self.assertTrue(ranges)
        self.assertFalse(any(r["chart_source"] for r in ranges))
        for cf_range in ranges:
            self.assertLess(cf_range["end_row"], cf_range["grand_row"])
            for subtotal_row in cf_range["subtotal_rows"]:
                self.assertFalse(
                    cf_range["start_row"] <= subtotal_row <= cf_range["end_row"],
                    cf_range,
                )

    def test_excel_conditional_formatting_uses_green_pass_fill_and_same_row_target(self):
        with tempfile.TemporaryDirectory() as tmp:
            output_path = tes.export_summary_excel(
                self.make_df(),
                {("REGION", "WEST"): 70},
                output_dir=tmp,
                cluster_name="Indulgence",
            )
            ws = load_workbook(output_path)["Noodle"]

        rules = []
        for cf_range, cf_rules in ws.conditional_formatting._cf_rules.items():
            for rule in cf_rules:
                color = getattr(rule.dxf.fill.fgColor, "rgb", "") if rule.dxf and rule.dxf.fill else ""
                rules.append((str(cf_range), rule.formula[0], color))

        pass_rules = [rule for rule in rules if ">=$" in rule[1]]
        fail_rules = [rule for rule in rules if "<$" in rule[1]]

        self.assertTrue(pass_rules)
        self.assertTrue(fail_rules)
        self.assertTrue(any(str(color).endswith("C4D79B") for _, _, color in pass_rules))
        self.assertTrue(any(str(color).endswith("E6B8B7") for _, _, color in fail_rules))
        cf_fonts = [
            rule.dxf.font
            for _cf_range, cf_rules in ws.conditional_formatting._cf_rules.items()
            for rule in cf_rules
            if rule.dxf
        ]
        self.assertTrue(all(font is None for font in cf_fonts))
        for _cell_range, formula, _color in rules:
            left, right = formula.split(">=") if ">=" in formula else formula.split("<")
            self.assertEqual(
                "".join(ch for ch in left if ch.isdigit()),
                "".join(ch for ch in right if ch.isdigit()),
            )

    def test_export_summary_excel_preserves_existing_embedded_target_rows(self):
        target_rows = [
            ["Division", "Dimension", "Name", "Target"],
            ["Noodle", "REGION", "WEST", 80],
            ["Snack", "CHANNEL", "GT", 72],
        ]
        targets = tes._target_rows_to_dict(target_rows)

        with tempfile.TemporaryDirectory() as tmp:
            output_path = tes.export_summary_excel(
                self.make_df(),
                targets,
                output_dir=tmp,
                cluster_name="Indulgence",
                target_rows=target_rows,
            )
            ws = load_workbook(output_path)["TARGETS"]
            rows = [
                [ws.cell(row=row_number, column=col).value for col in range(1, 5)]
                for row_number in range(1, ws.max_row + 1)
            ]

        self.assertEqual(rows, target_rows)

    def test_export_summary_excel_applies_targets_by_division_dimension_and_name(self):
        target_rows = [
            ["Division", "Dimension", "Name", "Target"],
            ["Noodle", "REGION", "WEST", 80],
            ["Snack", "REGION", "EAST", 55],
        ]
        targets = tes._target_rows_to_dict(target_rows)

        with tempfile.TemporaryDirectory() as tmp:
            output_path = tes.export_summary_excel(
                self.make_df(),
                targets,
                output_dir=tmp,
                cluster_name="Indulgence",
                target_rows=target_rows,
            )
            wb = load_workbook(output_path)
            noodle_rows = list(wb["Noodle"].iter_rows(values_only=True))
            snack_rows = list(wb["Snack"].iter_rows(values_only=True))

        noodle_region = next(row for row in noodle_rows if row[:2] == ("Noodle", "WEST"))
        snack_region = next(row for row in snack_rows if row[:2] == ("Snack", "EAST"))
        self.assertEqual(noodle_region[2], 80)
        self.assertEqual(snack_region[2], 55)

    def test_excel_category_charts_have_data_labels_and_non_overlapping_category_rows(self):
        df = pd.concat([
            self.make_df(),
            dataframe_with_produsen([
                {
                    "Source Division": "Noodle",
                    "Region": "WEST",
                    "Channel": "MT",
                    "Account": "ALPHA",
                    "Category Channel": "CUP NOODLE",
                    "Brand": "POP MIE",
                    "Period": "Jan 25",
                    "Store Code": "S003",
                    "Product Code": "SKU003",
                    "Facing": 12,
                },
                {
                    "Source Division": "Noodle",
                    "Region": "WEST",
                    "Channel": "MT",
                    "Account": "ALPHA",
                    "Category Channel": "CUP NOODLE",
                    "Brand": "COMP",
                    "Period": "Jan 25",
                    "Store Code": "S003",
                    "Product Code": "COMPETITOR SKU",
                    "Facing": 8,
                },
            ]),
        ], ignore_index=True)

        with tempfile.TemporaryDirectory() as tmp:
            output_path = tes.export_summary_excel(
                df,
                {("REGION", "WEST"): 70},
                output_dir=tmp,
                cluster_name="Indulgence",
            )
            ws = load_workbook(output_path)["Noodle"]

        category_charts = [chart for chart in ws._charts if " - " in str(chart.title)]
        anchors = sorted(chart.anchor._from.row for chart in category_charts)
        self.assertGreaterEqual(len(category_charts), 2)
        self.assertGreaterEqual(anchors[1] - anchors[0], tes.EXCEL_CATEGORY_CHART_ROW_STEP - 1)
        self.assertTrue(all(chart.dLbls and chart.dLbls.showVal for chart in category_charts))
        self.assertTrue(all(chart.x_axis.delete is False for chart in category_charts))
        self.assertTrue(all(chart.x_axis.tickLblPos == "low" for chart in category_charts))
        self.assertTrue(all(chart.series and chart.series[0].cat is not None for chart in category_charts))
        self.assertTrue(all(chart.series[0].cat.strRef is not None for chart in category_charts))

    def test_export_summary_excel_has_no_vba_dependency_or_macro_output(self):
        with open("tes.py", encoding="utf-8") as f:
            source = f.read()

        self.assertNotIn("win32com", source)
        self.assertNotIn("DispatchEx", source)
        self.assertNotIn("Worksheet_Change", source)
        self.assertNotIn("Workbook_Open", source)
        self.assertNotIn("add_vba_to_workbook", source)
        self.assertNotIn("enable_vba", source)

    def test_proses_data_exports_summary_xlsx_from_embedded_targets(self):
        original_cwd = os.getcwd()
        original_baca_semua_csv = tes.baca_semua_csv
        with tempfile.TemporaryDirectory() as tmp:
            try:
                os.chdir(tmp)
                tes.baca_semua_csv = self.make_df

                tes.proses_data()

                expected_name = f"Summary SOS_{os.path.basename(tmp)}.xlsx"
                output_path = os.path.join(tmp, expected_name)
                self.assertTrue(os.path.exists(output_path))
                self.assertFalse(os.path.exists(f"Summary SOS_{os.path.basename(tmp)}.xlsm"))
                self.assertEqual(load_workbook(output_path).sheetnames, ["Noodle", "Snack", "TARGETS"])
            finally:
                tes.baca_semua_csv = original_baca_semua_csv
                os.chdir(original_cwd)


class StoreDetailTests(unittest.TestCase):
    def test_period_is_pivot_axis_not_left_identity_column(self):
        df = dataframe_with_produsen([
            {
                "Period": "Jan 25",
                "Visit Date": pd.Timestamp("2025-01-02"),
                "Week": "W1",
                "Region": "WEST",
                "Area": "BANDUNG",
                "Channel": "MT",
                "Account": "ACC",
                "Store Code": "S001",
                "Store Name": "Store One",
                "Username": "user1",
                "Full Name": "User One",
                "Product Code": "SKU001",
                "Facing": 10,
            },
            {
                "Period": "Jan 25",
                "Visit Date": pd.Timestamp("2025-01-02"),
                "Week": "W1",
                "Region": "WEST",
                "Area": "BANDUNG",
                "Channel": "MT",
                "Account": "ACC",
                "Store Code": "S001",
                "Store Name": "Store One",
                "Username": "user1",
                "Full Name": "User One",
                "Product Code": "COMPETITOR SKU",
                "Facing": 5,
            },
        ])
        ws = FakeWorksheet()

        tes.buat_store_detail(ws, df)

        header1 = ws.updated_values[0]
        self.assertEqual(header1[:3], ["Region", "Area", "Store Name"])
        self.assertNotIn("Period", header1[:3])
        self.assertIn("Jan 25", header1)
        self.assertEqual(ws.frozen, (2, 3))


class DashboardDivisionAwareTests(unittest.TestCase):
    def setUp(self):
        self._orig_sleep = tes.time.sleep
        tes.time.sleep = lambda _seconds: None

    def tearDown(self):
        tes.time.sleep = self._orig_sleep

    def make_dashboard_df(self):
        return dataframe_with_produsen([
            {
                "Source Division": "Noodle",
                "Region": "WEST",
                "Channel": "MT",
                "Account": "ALPHA",
                "Category Channel": "BAG NOODLE",
                "Brand": "INDOMIE",
                "Period": "Jan 25",
                "Store Code": "S001",
                "Product Code": "SKU001",
                "Facing": 30,
            },
            {
                "Source Division": "Noodle",
                "Region": "WEST",
                "Channel": "MT",
                "Account": "ALPHA",
                "Category Channel": "BAG NOODLE",
                "Brand": "COMP",
                "Period": "Jan 25",
                "Store Code": "S001",
                "Product Code": "COMPETITOR SKU",
                "Facing": 10,
            },
            {
                "Source Division": "Snack",
                "Region": "EAST",
                "Channel": "GT",
                "Account": "BETA",
                "Category Channel": "OTHER",
                "Brand": "CHIKI",
                "Period": "Jan 25",
                "Store Code": "S002",
                "Product Code": "SKU002",
                "Facing": 10,
            },
            {
                "Source Division": "Snack",
                "Region": "EAST",
                "Channel": "GT",
                "Account": "BETA",
                "Category Channel": "OTHER",
                "Brand": "COMP",
                "Period": "Jan 25",
                "Store Code": "S002",
                "Product Code": "COMPETITOR SKU",
                "Facing": 30,
            },
        ])

    def test_monthly_sos_table_can_group_by_source_division_and_region(self):
        df = dataframe_with_produsen([
            {
                "Source Division": "Noodle",
                "Region": "WEST",
                "Period": "Jan 25",
                "Product Code": "SKU001",
                "Facing": 30,
            },
            {
                "Source Division": "Noodle",
                "Region": "WEST",
                "Period": "Jan 25",
                "Product Code": "COMPETITOR SKU",
                "Facing": 10,
            },
            {
                "Source Division": "Snack",
                "Region": "WEST",
                "Period": "Jan 25",
                "Product Code": "SKU002",
                "Facing": 10,
            },
            {
                "Source Division": "Snack",
                "Region": "WEST",
                "Period": "Jan 25",
                "Product Code": "COMPETITOR SKU",
                "Facing": 30,
            },
        ])

        rows, meta = tes.buat_tabel_sos_monthly(
            df,
            ["Source Division", "Region"],
            "REGION",
            ["Jan 25"],
            {("REGION", "WEST"): 70},
        )

        self.assertEqual(rows[0][:3], ["DIVISION", "Region", "TARGET"])
        self.assertIn(["Noodle", "WEST", 70, 30, 10, 40, 75.0], rows)
        self.assertIn(["Snack", "WEST", 70, 10, 30, 40, 25.0], rows)
        self.assertIn(["Noodle TOTAL", "", 65.0, 30, 10, 40, 75.0], rows)
        self.assertIn(["Snack TOTAL", "", 65.0, 10, 30, 40, 25.0], rows)
        self.assertEqual(meta["target_col_idx"], 2)
        self.assertEqual(meta["sos_col_indices"], [6])

    def test_channel_account_table_keeps_source_divisions_separate(self):
        df = dataframe_with_produsen([
            {
                "Source Division": "Noodle",
                "Channel": "MT",
                "Account": "ALPHA",
                "Period": "Jan 25",
                "Store Code": "S001",
                "Product Code": "SKU001",
                "Facing": 8,
            },
            {
                "Source Division": "Noodle",
                "Channel": "MT",
                "Account": "ALPHA",
                "Period": "Jan 25",
                "Store Code": "S001",
                "Product Code": "COMPETITOR SKU",
                "Facing": 2,
            },
            {
                "Source Division": "Snack",
                "Channel": "MT",
                "Account": "ALPHA",
                "Period": "Jan 25",
                "Store Code": "S002",
                "Product Code": "SKU002",
                "Facing": 3,
            },
            {
                "Source Division": "Snack",
                "Channel": "MT",
                "Account": "ALPHA",
                "Period": "Jan 25",
                "Store Code": "S002",
                "Product Code": "COMPETITOR SKU",
                "Facing": 7,
            },
        ])

        rows, meta = tes.buat_tabel_channel_account(
            df,
            ["Jan 25"],
            {("CHANNEL-ACCOUNT", "MT - ALPHA"): 65},
        )

        self.assertEqual(rows[0][:4], ["DIVISION", "CHANNEL", "ACCOUNT", "TARGET"])
        self.assertIn(["Noodle", "MT", "ALPHA", 65, 8, 2, 10, 80.0], rows)
        self.assertIn(["Snack", "MT", "ALPHA", 65, 3, 7, 10, 30.0], rows)
        self.assertIn(["Noodle", "MT TOTAL", "", "", 8, 2, 10, 80.0], rows)
        self.assertIn(["Snack", "MT TOTAL", "", "", 3, 7, 10, 30.0], rows)
        self.assertEqual(meta["target_col_idx"], 3)
        self.assertEqual(meta["sos_col_indices"], [7])

    def test_dashboard_region_section_uses_division_dimension_when_available(self):
        df = self.make_dashboard_df()
        ws = FakeWorksheet()

        tes.buat_dashboard(ws, df)

        title_idx = ws.updated_values.index(["SOS% BY REGION"])
        self.assertEqual(
            ws.updated_values[title_idx + 1][:3],
            ["DIVISION", "Region", "TARGET"],
        )

        titles = [row[0] for row in ws.updated_values if row]
        self.assertNotIn("SOS% BY REGION x DIVISI", titles)
        self.assertNotIn("SOS% BY ACCOUNT x DIVISI", titles)
        self.assertNotIn("SOS% BY CATEGORY BY DIVISI", titles)

    def test_dashboard_all_mode_preserves_selector_and_shows_full_division_summary(self):
        ws = FakeWorksheet(b1_value="ALL")

        tes.buat_dashboard(ws, self.make_dashboard_df())

        self.assertFalse(ws.cleared)
        self.assertIn(["A2:ZZZ"], ws.batch_clears)
        self.assertEqual(ws.updated_values[0], ["DIVISION", "ALL"])
        self.assertEqual(ws.updated_values[2], ["DIVISION SUMMARY"])
        self.assertIn(["Noodle", 75.0], ws.updated_values)
        self.assertIn(["Snack", 25.0], ws.updated_values)

    def test_dashboard_single_division_mode_keeps_all_rows_for_apps_script_filter(self):
        ws = FakeWorksheet(b1_value="Noodle")

        tes.buat_dashboard(ws, self.make_dashboard_df())

        self.assertEqual(ws.updated_values[0], ["DIVISION", "Noodle"])
        self.assertIn(["Noodle", 75.0], ws.updated_values)
        self.assertIn(["Snack", 25.0], ws.updated_values)

        region_title = ws.updated_values.index(["SOS% BY REGION"])
        next_title = ws.updated_values.index(["SOS% BY CHANNEL"])
        region_rows = ws.updated_values[region_title:next_title]
        self.assertIn(["Noodle", "WEST", 65.0, 30, 10, 40, 75.0], region_rows)
        self.assertIn(["Snack", "EAST", 65.0, 10, 30, 40, 25.0], region_rows)

    def test_dashboard_invalid_division_selector_falls_back_to_all(self):
        ws = FakeWorksheet(b1_value="Bad Division")

        tes.buat_dashboard(ws, self.make_dashboard_df())

        self.assertEqual(ws.updated_values[0], ["DIVISION", "ALL"])
        region_title = ws.updated_values.index(["SOS% BY REGION"])
        next_title = ws.updated_values.index(["SOS% BY CHANNEL"])
        region_rows = ws.updated_values[region_title:next_title]
        self.assertIn(["Noodle", "WEST", 65.0, 30, 10, 40, 75.0], region_rows)
        self.assertIn(["Snack", "EAST", 65.0, 10, 30, 40, 25.0], region_rows)

    def chart_requests_from(self, ws):
        chart_requests = []
        for args, _kwargs in ws.spreadsheet.batch_updates:
            if args and isinstance(args[0], dict):
                chart_requests.extend([
                    r for r in args[0].get("requests", [])
                    if "addChart" in r
                ])
        return chart_requests

    def test_dashboard_all_mode_creates_division_summary_chart(self):
        ws = FakeWorksheet(b1_value="ALL")

        tes.buat_dashboard(ws, self.make_dashboard_df())

        chart_requests = self.chart_requests_from(ws)
        titles = [
            r["addChart"]["chart"]["spec"].get("title")
            for r in chart_requests
        ]
        self.assertEqual(titles, ["DIVISION SUMMARY"])
        row_index = chart_requests[0]["addChart"]["chart"]["position"]["overlayPosition"]["anchorCell"]["rowIndex"]
        self.assertGreaterEqual(row_index, len(ws.updated_values))

    def test_dashboard_selected_division_reuses_category_by_divisi_chart(self):
        ws = FakeWorksheet(b1_value="Noodle")

        tes.buat_dashboard(ws, self.make_dashboard_df())

        chart_requests = self.chart_requests_from(ws)
        titles = [
            r["addChart"]["chart"]["spec"].get("title", "")
            for r in chart_requests
        ]
        self.assertTrue(any("BAG NOODLE" in title for title in titles))
        self.assertNotIn("DIVISION SUMMARY", titles)
        row_indexes = [
            r["addChart"]["chart"]["position"]["overlayPosition"]["anchorCell"]["rowIndex"]
            for r in chart_requests
        ]
        self.assertTrue(all(row_index >= len(ws.updated_values) for row_index in row_indexes))


if __name__ == "__main__":
    unittest.main()
