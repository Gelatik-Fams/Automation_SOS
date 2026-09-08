"""Target parsing, defaults, enrichment, and Summary workbook handling."""

from openpyxl import load_workbook
import os

from sos_engine.common import (
    DEFAULT_TARGET,
    TARGET_DIM_ORDER,
    filter_dashboard_by_division,
    get_source_divisions,
)


def get_target(targets, dim, nama):
    """Cari target SOS% — fallback ke DEFAULT lalu 65."""
    dim_up  = dim.upper()
    nama_up = str(nama).upper()
    return targets.get((dim_up, nama_up),
           targets.get((dim_up, 'DEFAULT'), 65.0))


def default_target_rows(divisions=None):
    divisions = list(divisions or ['ALL'])
    rows = [['Division', 'Dimension', 'Name', 'Target']]
    for division in divisions:
        rows.extend([[division, dim, 'DEFAULT', DEFAULT_TARGET] for dim in TARGET_DIM_ORDER])
    return rows


def default_targets_dict():
    return {
        (dim, 'DEFAULT'): DEFAULT_TARGET
        for dim in TARGET_DIM_ORDER
    }


def _parse_target_value(value, row_number):
    if value is None or str(value).strip() == '':
        raise ValueError(f'TARGETS row {row_number}: nilai target kosong.')
    try:
        if isinstance(value, str):
            value = value.replace('%', '').replace(',', '.').strip()
        return float(value)
    except ValueError as exc:
        raise ValueError(f'TARGETS row {row_number}: nilai target tidak valid ({value}).') from exc


def _target_rows_to_dict(target_rows):
    if not target_rows:
        raise ValueError('TARGETS tidak valid: tidak ada data target.')

    header = [str(v or '').strip().upper() for v in target_rows[0]]
    if header[:4] == ['DIVISION', 'DIMENSION', 'NAME', 'TARGET']:
        dim_idx, name_idx, target_idx = 1, 2, 3
    elif header[:3] == ['DIMENSION', 'NAME', 'TARGET (%)']:
        dim_idx, name_idx, target_idx = 0, 1, 2
    else:
        raise ValueError('TARGETS tidak valid: header harus Division | Dimension | Name | Target.')

    targets = {}
    for row_number, row in enumerate(target_rows[1:], start=2):
        values = list(row) + [''] * 4
        dim = str(values[dim_idx] or '').strip().upper()
        nama = str(values[name_idx] or '').strip().upper()
        target_value = values[target_idx]

        if not dim and not nama and (target_value is None or str(target_value).strip() == ''):
            continue
        if not dim or not nama:
            raise ValueError(f'TARGETS row {row_number}: Dimension dan Name wajib diisi.')

        targets[(dim, nama)] = _parse_target_value(target_value, row_number)

    if not targets:
        raise ValueError('TARGETS tidak valid: tidak ada data target.')

    return targets


def _build_division_targets(target_rows, division):
    """Build a (dim, name) → value dict scoped to one division.

    Merges ALL-division rows (base) with division-specific rows (override),
    so per-division values take precedence while global defaults still apply.
    """
    if not target_rows:
        return default_targets_dict()
    header = target_rows[0]
    all_data = [r for r in target_rows[1:] if str(r[0] or '').strip().upper() == 'ALL']
    div_data = [r for r in target_rows[1:] if r[0] == division]
    merged = [header] + all_data + div_data
    if len(merged) == 1:
        merged = target_rows
    return _target_rows_to_dict(merged)


def _enrich_targets_with_df_values(df, target_rows):
    """Append-only: tambah baris TARGETS untuk setiap nilai dimensi di df yang belum ada.

    Nilai awal tiap baris baru = DEFAULT target untuk division+dim tersebut.
    Baris yang sudah ada tidak pernah diubah.
    """
    DIM_COL_MAP = [
        ('REGION',           'Region'),
        ('CHANNEL',          'Channel'),
        ('ACCOUNT GELATIK',  'Account'),
        ('CATEGORY CHANNEL', 'Category Channel'),
    ]

    base = target_rows if target_rows else [['Division', 'Dimension', 'Name', 'Target']]

    existing = set()
    for row in base[1:]:
        if len(row) >= 3 and row[0] is not None:
            existing.add((str(row[0]).strip(), str(row[1]).strip().upper(), str(row[2]).strip().upper()))

    def _seed(division, dim):
        for scope in (division, 'ALL'):
            for row in base[1:]:
                if (len(row) >= 4
                        and str(row[0]).strip() == scope
                        and str(row[1]).strip().upper() == dim
                        and str(row[2]).strip().upper() == 'DEFAULT'):
                    try:
                        return float(row[3])
                    except (ValueError, TypeError):
                        pass
        return DEFAULT_TARGET

    new_rows = []
    for division in get_source_divisions(df):
        df_div = filter_dashboard_by_division(df, division)

        for dim, col in DIM_COL_MAP:
            if col not in df_div.columns:
                continue
            seed = _seed(division, dim)
            for val in sorted(str(v).strip().upper() for v in df_div[col].dropna().unique() if str(v).strip()):
                if (division, dim, val) not in existing:
                    new_rows.append([division, dim, val, seed])
                    existing.add((division, dim, val))

        if 'Channel' in df_div.columns and 'Account' in df_div.columns:
            pairs = (df_div[['Channel', 'Account']]
                     .drop_duplicates()
                     .dropna()
                     .sort_values(['Channel', 'Account']))
            seed_ca = _seed(division, 'CHANNEL-ACCOUNT')
            for _, r in pairs.iterrows():
                ch, acc = str(r['Channel']).strip(), str(r['Account']).strip()
                if not ch or not acc:
                    continue
                val = f'{ch} - {acc}'.upper()
                if (division, 'CHANNEL-ACCOUNT', val) not in existing:
                    new_rows.append([division, 'CHANNEL-ACCOUNT', val, seed_ca])
                    existing.add((division, 'CHANNEL-ACCOUNT', val))

        # Add GRAND TOTAL entry for each dimension so users can edit it in TARGETS
        all_dims = [dim for dim, _ in DIM_COL_MAP] + ['CHANNEL-ACCOUNT']
        for dim in all_dims:
            gt_key = (division, dim, 'GRAND TOTAL')
            if gt_key not in existing:
                seed = _seed(division, dim)
                new_rows.append([division, dim, 'GRAND TOTAL', seed])
                existing.add(gt_key)

    if not new_rows:
        return target_rows
    return base + new_rows


def read_targets_from_workbook(path):
    wb = None
    try:
        wb = load_workbook(path, data_only=True)
    except Exception as exc:
        raise ValueError(f'Workbook summary ada tetapi tidak bisa dibaca: {exc}') from exc

    try:
        if 'TARGETS' not in wb.sheetnames:
            raise ValueError('Workbook summary tidak valid: sheet TARGETS tidak ditemukan.')

        ws = wb['TARGETS']
        max_col = max(ws.max_column, 4)
        raw_rows = [
            [ws.cell(row=row_number, column=col).value for col in range(1, max_col + 1)]
            for row_number in range(1, ws.max_row + 1)
        ]

        header = [str(v or '').strip().upper() for v in raw_rows[0]]
        if header[:4] == ['DIVISION', 'DIMENSION', 'NAME', 'TARGET']:
            target_rows = [row[:4] for row in raw_rows]
        elif header[:3] == ['DIMENSION', 'NAME', 'TARGET (%)']:
            target_rows = [['Division', 'Dimension', 'Name', 'Target']]
            for row in raw_rows[1:]:
                values = list(row) + [''] * 3
                target_rows.append(['ALL', values[0], values[1], values[2]])
        else:
            raise ValueError('TARGETS tidak valid: header harus Division | Dimension | Name | Target.')
    finally:
        wb.close()

    return _target_rows_to_dict(target_rows), target_rows


def load_or_create_summary_targets(summary_path):
    if not os.path.exists(summary_path):
        return default_targets_dict(), None
    return read_targets_from_workbook(summary_path)
